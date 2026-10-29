import csv
import os
import io
import threading
import random
from datetime import datetime, timedelta
from flask import Flask, request, jsonify, render_template_string, send_file

# ==========================================
# CONFIGURATION
# ==========================================
FICHIER_RESULTATS = "resultats_backyard.csv"
FICHIER_PARTICIPANTS = "participants.csv"
DUREE_BOUCLE_MINUTES = 60
DELAI_DE_GRACE_SECONDES = 60
DISTANCE_BOUCLE_KM = 6.7 
CODE_PIN_ADMIN = "1234"

app = Flask(__name__)
lock = threading.Lock()

# --- VARIABLES GLOBALES EN MÉMOIRE ---
etat_course = {
    "start_time": None,
    "historique_boucles": {},
    "coureurs": {},
    "logs_recents": [],
    "temps_coureurs": {},
    "dnf": set()
}

def charger_participants():
    if not os.path.exists(FICHIER_PARTICIPANTS):
        with open(FICHIER_PARTICIPANTS, mode="w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            writer.writerow(["Dossard", "Nom", "Prenom"])
            writer.writerow(["42", "Exemple", "Coureur"])
    else:
        with open(FICHIER_PARTICIPANTS, mode="r", encoding="utf-8") as f:
            reader = csv.reader(f, delimiter=';') 
            next(reader, None)
            for row in reader:
                if len(row) >= 3:
                    etat_course["coureurs"][row[0].strip()] = f"{row[2].strip()} {row[1].strip()}"

def reconstruire_etat():
    global etat_course
    etat_course["start_time"] = None
    etat_course["historique_boucles"] = {}
    etat_course["temps_coureurs"] = {}
    etat_course["dnf"] = set()
    etat_course["logs_recents"] = []
    
    if not os.path.exists(FICHIER_RESULTATS):
        return
        
    start_time_calcule = None
    lignes_valides = []
    
    try:
        with open(FICHIER_RESULTATS, mode="r", encoding="utf-8") as f:
            reader = csv.reader(f)
            header = next(reader, None)
            for row in reader:
                if len(row) >= 7:
                    lignes_valides.append(row)
                    
        for row in lignes_valides:
            b = int(row[0])
            d = row[1]
            h_arr_str = row[3]
            t_boucle_str = row[4]
            statut = row[6]
            
            if b not in etat_course["historique_boucles"]:
                etat_course["historique_boucles"][b] = set()
                etat_course["temps_coureurs"][b] = {}
            
            if not start_time_calcule and t_boucle_str != "--:--":
                try:
                    t_parts = t_boucle_str.split(":")
                    t_m, t_s = int(t_parts[0]), int(t_parts[1])
                    
                    h_parts = h_arr_str.split(":")
                    h_h, h_m, h_s = int(h_parts[0]), int(h_parts[1]), int(h_parts[2])
                    
                    maintenant = datetime.now()
                    arrivee_dt = maintenant.replace(hour=h_h, minute=h_m, second=h_s, microsecond=0)
                    
                    if arrivee_dt > maintenant + timedelta(hours=2):
                        arrivee_dt -= timedelta(days=1)
                    
                    delta_boucle = timedelta(minutes=t_m, seconds=t_s)
                    start_time_calcule = arrivee_dt - delta_boucle - timedelta(minutes=(b-1)*DUREE_BOUCLE_MINUTES)
                except:
                    pass

            if "OK" in statut or "In extremis" in statut:
                etat_course["historique_boucles"][b].add(d)
            
            if "DNF" in statut:
                etat_course["dnf"].add(d)
            
            try:
                t_parts = t_boucle_str.split(":")
                sec = int(t_parts[0]) * 60 + int(t_parts[1])
                etat_course["temps_coureurs"][b][d] = {"str": t_boucle_str, "sec": sec}
            except:
                pass
                
        if start_time_calcule:
            etat_course["start_time"] = start_time_calcule
            etat_course["logs_recents"].insert(0, {"text": f"Course restaurée avec succès ({len(lignes_valides)} passages).", "color": "#89b4fa"})
            
    except Exception as e:
        print("Erreur restauration:", e)

def initialiser_fichiers_et_etat():
    charger_participants()
    if os.path.exists(FICHIER_RESULTATS):
        with open(FICHIER_RESULTATS, mode="r", encoding="utf-8") as f:
            lignes = f.readlines()
        if len(lignes) > 1:
            reconstruire_etat()
            return

    if not os.path.exists(FICHIER_RESULTATS):
        with open(FICHIER_RESULTATS, mode="w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            writer.writerow(["Boucle", "Dossard", "Nom", "Heure_Arrivee", "Temps_Boucle", "Temps_Repos", "Statut"])

initialiser_fichiers_et_etat()

def obtenir_etat_boucle():
    if not etat_course["start_time"]:
        return 1, None, None
    maintenant = datetime.now()
    secondes_ecoulees = (maintenant - etat_course["start_time"]).total_seconds()
    num_boucle = int(secondes_ecoulees // (DUREE_BOUCLE_MINUTES * 60)) + 1
    debut = etat_course["start_time"] + timedelta(minutes=(num_boucle - 1) * DUREE_BOUCLE_MINUTES)
    fin = debut + timedelta(minutes=DUREE_BOUCLE_MINUTES)
    return num_boucle, debut, fin

def sauvegarder_statuts_fin_boucle(derniere_boucle_terminee):
    try:
        import openpyxl
        from openpyxl.styles import Font, Alignment, PatternFill
    except ImportError:
        return

    wb = openpyxl.Workbook()
    
    font_header = Font(bold=True, color="FFFFFF")
    fill_header = PatternFill(start_color="313244", end_color="313244", fill_type="solid")
    fill_en_lice = PatternFill(start_color="A6E3A1", end_color="A6E3A1", fill_type="solid")
    fill_dnf = PatternFill(start_color="F38BA8", end_color="F38BA8", fill_type="solid")
    align_center = Alignment(horizontal="center", vertical="center")
    
    dnf_loop_map = {}
    for d in etat_course["coureurs"]:
        if d in etat_course["dnf"] or d not in etat_course["historique_boucles"].get(derniere_boucle_terminee, set()):
            loop_missed = 1
            for b in range(1, derniere_boucle_terminee + 1):
                if d not in etat_course["historique_boucles"].get(b, set()):
                    loop_missed = b
                    break
            dnf_loop_map[d] = loop_missed

    def sort_dossard(d_id):
        try: return int(d_id)
        except ValueError: return 999999

    coureurs_tries = sorted(etat_course["coureurs"].keys(), key=sort_dossard)

    ws_global = wb.active
    ws_global.title = "État Global Actuel"
    ws_global.append(["Dossard", "Nom", "Statut Actuel", "Boucle d'élimination"])
    
    for col in range(1, 5):
        cell = ws_global.cell(row=1, column=col)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center

    row_idx = 2
    for d in coureurs_tries:
        nom = etat_course["coureurs"][d]
        if d in dnf_loop_map and dnf_loop_map[d] <= derniere_boucle_terminee:
            statut = "DNF"
            b_elim = f"Boucle {dnf_loop_map[d]}"
            fill_color = fill_dnf
        else:
            statut = "En lice"
            b_elim = "-"
            fill_color = fill_en_lice

        ws_global.append([d, nom, statut, b_elim])
        for col in range(1, 5):
            ws_global.cell(row=row_idx, column=col).alignment = align_center
        ws_global.cell(row=row_idx, column=3).fill = fill_color
        row_idx += 1

    for b in range(1, derniere_boucle_terminee + 1):
        ws_b = wb.create_sheet(title=f"Fin Boucle {b}")
        ws_b.append(["Dossard", "Nom", "Statut à la fin de cette boucle"])
        
        for col in range(1, 4):
            cell = ws_b.cell(row=1, column=col)
            cell.font = font_header
            cell.fill = fill_header
            cell.alignment = align_center
            
        row_idx_b = 2
        for d in coureurs_tries:
            nom = etat_course["coureurs"][d]
            
            if d in etat_course["historique_boucles"].get(b, set()):
                statut = "En lice"
                fill_color = fill_en_lice
            elif d in dnf_loop_map and dnf_loop_map[d] == b:
                statut = "Nouveau DNF"
                fill_color = fill_dnf
            else:
                statut = "Déjà DNF"
                fill_color = PatternFill(start_color="E6E9EF", end_color="E6E9EF", fill_type="solid")
            
            ws_b.append([d, nom, statut])
            for col in range(1, 4):
                ws_b.cell(row=row_idx_b, column=col).alignment = align_center
            ws_b.cell(row=row_idx_b, column=3).fill = fill_color
            row_idx_b += 1
            
        ws_b.column_dimensions['A'].width = 15
        ws_b.column_dimensions['B'].width = 30
        ws_b.column_dimensions['C'].width = 35

    ws_global.column_dimensions['A'].width = 15
    ws_global.column_dimensions['B'].width = 30
    ws_global.column_dimensions['C'].width = 20
    ws_global.column_dimensions['D'].width = 25

    wb.save("Suivi_Statuts_Coureurs.xlsx")

# ==========================================
# ROUTES API
# ==========================================

@app.route("/api/restore", methods=["POST"])
def restore_race():
    with lock:
        if 'file' not in request.files:
            return jsonify({"status": "error", "message": "Aucun fichier reçu."})
        file = request.files['file']
        if not file.filename.endswith('.csv'):
            return jsonify({"status": "error", "message": "Format invalide (.csv attendu)."})
            
        if os.path.exists(FICHIER_RESULTATS):
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            os.rename(FICHIER_RESULTATS, f"resultats_ecrase_avant_restore_{timestamp}.csv")
            
        file.save(FICHIER_RESULTATS)
        reconstruire_etat()
        return jsonify({"status": "ok"})

@app.route("/api/verify_pin", methods=["POST"])
def verify_pin():
    pin_soumis = request.json.get("pin", "")
    if pin_soumis == CODE_PIN_ADMIN:
        return jsonify({"status": "ok"})
    return jsonify({"status": "error", "message": "Code PIN incorrect."})

@app.route("/api/participants", methods=["GET"])
def get_participants():
    if not etat_course["start_time"]:
        survivants = set(etat_course["coureurs"].keys())
    else:
        num_boucle, _, _ = obtenir_etat_boucle()
        survivants = set(etat_course["coureurs"].keys())
        for b in range(1, num_boucle):
            survivants = survivants.intersection(etat_course["historique_boucles"].get(b, set()))
        survivants = survivants - etat_course["dnf"]

    liste = [{"dossard": d, "nom": nom} for d, nom in etat_course["coureurs"].items() if d in survivants]
    return jsonify({"participants": liste})

@app.route("/api/state", methods=["GET"])
def get_state():
    if not etat_course["start_time"]:
        return jsonify({"started": False, "logs": etat_course["logs_recents"]})

    num_boucle, debut_boucle, fin_boucle = obtenir_etat_boucle()
    maintenant = datetime.now()
    temps_restant = int((fin_boucle - maintenant).total_seconds())
    
    if num_boucle not in etat_course["historique_boucles"]:
        if num_boucle > 1:
            sauvegarder_statuts_fin_boucle(num_boucle - 1)
            
        etat_course["historique_boucles"][num_boucle] = set()
        etat_course["temps_coureurs"][num_boucle] = {}
        etat_course["logs_recents"].insert(0, {"text": f"--- DÉPART BOUCLE {num_boucle} ({debut_boucle.strftime('%H:%M')}) ---", "color": "#f9e2af"})

    survivants = set(etat_course["coureurs"].keys())
    for b in range(1, num_boucle):
        survivants = survivants.intersection(etat_course["historique_boucles"].get(b, set()))
    survivants = survivants - etat_course["dnf"]

    runners_data = []
    for d, nom in etat_course["coureurs"].items():
        if d in etat_course["dnf"] or d not in survivants:
            status_code = "dnf"
        elif d in etat_course["historique_boucles"].get(num_boucle, set()):
            status_code = "arrive"
        else:
            status_code = "encourse"

        total_sec = 0
        nb_boucles = 0
        last_loop_str = "--:--"
        last_loop_sec = 999999
        fastest_sec = 999999
        fastest_boucle = 0
        
        for b in range(1, num_boucle + 1):
            if d in etat_course["temps_coureurs"].get(b, {}):
                sec = etat_course["temps_coureurs"][b][d]["sec"]
                total_sec += sec
                last_loop_str = etat_course["temps_coureurs"][b][d]["str"]
                last_loop_sec = sec
                nb_boucles += 1
                if sec < fastest_sec:
                    fastest_sec = sec
                    fastest_boucle = b
                    
        just_beat_record = False
        if status_code == "arrive" and fastest_boucle == num_boucle and nb_boucles > 1:
            just_beat_record = True
                
        if nb_boucles > 0:
            moy_sec = total_sec // nb_boucles
            moy_str = f"{moy_sec//60:02d}:{moy_sec%60:02d}"
            allure_sec = int(moy_sec / DISTANCE_BOUCLE_KM)
            allure_str = f"{allure_sec//60:02d}:{allure_sec%60:02d}/k"
        else:
            moy_str = "--:--"
            moy_sec = 999999
            allure_str = "--:--"

        runners_data.append({
            "dossard": d, "nom": nom, "status": status_code,
            "last_loop": last_loop_str, "last_loop_sec": last_loop_sec,
            "just_beat_record": just_beat_record, 
            "moyenne": moy_str, "moyenne_sec": moy_sec,
            "allure": allure_str, 
            "tours_valides": nb_boucles
        })

    # NOUVEAU TRI TV/SPEAKER : 1. Boucles valides (décroissant), 2. Temps moyen (croissant)
    runners_data.sort(key=lambda x: (-x["tours_valides"], x["moyenne_sec"]))

    record_absolu = {"str": "--:--", "nom": "", "dossard": ""}
    meilleur_temps = 999999
    for b_idx, coureurs_boucle in etat_course["temps_coureurs"].items():
        for d, data_temps in coureurs_boucle.items():
            if data_temps["sec"] < meilleur_temps:
                meilleur_temps = data_temps["sec"]
                record_absolu = {
                    "str": data_temps["str"], "dossard": d, "nom": etat_course["coureurs"].get(d, "Inconnu")
                }

    return jsonify({
        "started": True, "boucle": num_boucle, "debut_str": debut_boucle.strftime('%H:%M:%S'),
        "temps_restant": temps_restant, "survivants_total": len(survivants),
        "runners_all": runners_data, "record_absolu": record_absolu,
        "logs": etat_course["logs_recents"] 
    })

@app.route("/api/start", methods=["POST"])
def start_race():
    with lock:
        if not etat_course["start_time"]:
            etat_course["start_time"] = datetime.now()
            etat_course["historique_boucles"][1] = set()
            etat_course["temps_coureurs"][1] = {}
            return jsonify({"status": "ok"})
        return jsonify({"status": "already_started"})

@app.route("/api/submit", methods=["POST"])
def submit_dossard():
    dossard = request.json.get("dossard", "").strip()
    is_correction = request.json.get("correction", False)
    if not etat_course["start_time"] or not dossard: return jsonify({"status": "error", "message": "Erreur."})

    with lock:
        if dossard not in etat_course["coureurs"]: 
            return jsonify({"status": "error", "message": f"Le dossard {dossard} n'est pas dans la liste des participants !"})
            
        maintenant = datetime.now()
        num_boucle, debut_boucle, fin_boucle = obtenir_etat_boucle()
        secondes_depuis_debut = (maintenant - debut_boucle).total_seconds()
        
        if dossard in etat_course["dnf"]:
            return jsonify({"status": "error", "message": f"Le coureur {dossard} est déjà éliminé de la course."})

        survivants = set(etat_course["coureurs"].keys())
        for b in range(1, num_boucle):
            survivants = survivants.intersection(etat_course["historique_boucles"].get(b, set()))
        survivants = survivants - etat_course["dnf"]

        boucle_cible = num_boucle
        in_extremis = False

        if not is_correction:
            if dossard not in survivants:
                eligible_grace = True
                for b in range(1, num_boucle - 1):
                    if dossard not in etat_course["historique_boucles"].get(b, set()):
                        eligible_grace = False
                        break

                if eligible_grace and num_boucle > 1 and dossard not in etat_course["historique_boucles"].get(num_boucle - 1, set()) and secondes_depuis_debut <= DELAI_DE_GRACE_SECONDES:
                    boucle_cible = num_boucle - 1
                    in_extremis = True
                else:
                    return jsonify({"status": "error", "message": f"Le coureur {dossard} n'est plus en lice (boucle ratée précédemment)."})
        else:
            if num_boucle <= 1: return jsonify({"status": "error", "message": "Aucune boucle précédente."})
            
            eligible_correction = True
            for b in range(1, num_boucle - 1):
                if dossard not in etat_course["historique_boucles"].get(b, set()):
                    eligible_correction = False
                    break
                    
            if not eligible_correction:
                return jsonify({"status": "error", "message": f"Correction impossible : {dossard} avait déjà raté une boucle avant."})
                
            boucle_cible = num_boucle - 1
            in_extremis = True

        return valider_coureur(dossard, boucle_cible, in_extremis)

def valider_coureur(dossard, num_boucle, in_extremis):
    debut_cette_boucle = etat_course["start_time"] + timedelta(minutes=(num_boucle - 1) * DUREE_BOUCLE_MINUTES)
    fin_cette_boucle = debut_cette_boucle + timedelta(minutes=DUREE_BOUCLE_MINUTES)

    if num_boucle not in etat_course["historique_boucles"]:
        etat_course["historique_boucles"][num_boucle] = set()
        etat_course["temps_coureurs"][num_boucle] = {}

    if dossard in etat_course["historique_boucles"][num_boucle]:
        return jsonify({"status": "duplicate", "message": f"Le dossard {dossard} a déjà validé cette boucle !"})

    nom_coureur = etat_course["coureurs"].get(dossard, "Inconnu")
    
    heure_arrivee = (fin_cette_boucle - timedelta(seconds=1)) if in_extremis else datetime.now()
    statut = "OK (In extremis)" if in_extremis else "OK"
    temps_boucle = heure_arrivee - debut_cette_boucle
    temps_repos = fin_cette_boucle - heure_arrivee

    total_sec = int(temps_boucle.total_seconds())
    b_min, b_sec = divmod(total_sec, 60)
    str_temps_boucle = f"{b_min:02d}:{b_sec:02d}"
    
    etat_course["temps_coureurs"][num_boucle][dossard] = {"str": str_temps_boucle, "sec": total_sec}

    if temps_repos.total_seconds() < 0:
        statut = "DNF"
        str_temps_repos = "00:00"
        couleur = "#f38ba8" 
        etat_course["dnf"].add(dossard) 
    else:
        r_min, r_sec = divmod(int(temps_repos.total_seconds()), 60)
        str_temps_repos = f"{r_min:02d}:{r_sec:02d}"
        couleur = "#fab387" if in_extremis else "#a6e3a1" 

    with open(FICHIER_RESULTATS, mode="a", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow([num_boucle, dossard, nom_coureur, heure_arrivee.strftime("%H:%M:%S"), str_temps_boucle, str_temps_repos, statut])

    etat_course["historique_boucles"][num_boucle].add(dossard)
    etat_course["logs_recents"].insert(0, {"text": f"[B{num_boucle}] {dossard} - {nom_coureur} | Tour: {str_temps_boucle}", "color": couleur})
    return jsonify({"status": "ok"})

@app.route("/api/undo", methods=["POST"])
def undo_dossard():
    dossard = request.json.get("dossard", "").strip()
    if not etat_course["start_time"] or not dossard: return jsonify({"status": "error", "message": "Erreur."})

    with lock:
        num_boucle, _, _ = obtenir_etat_boucle()
        if dossard in etat_course["historique_boucles"].get(num_boucle, set()):
            etat_course["historique_boucles"][num_boucle].remove(dossard)
            if dossard in etat_course["temps_coureurs"].get(num_boucle, {}):
                del etat_course["temps_coureurs"][num_boucle][dossard]
            if dossard in etat_course["dnf"]:
                etat_course["dnf"].remove(dossard)
            
            lignes = []
            with open(FICHIER_RESULTATS, mode="r", encoding="utf-8") as f:
                reader = csv.reader(f)
                for row in reader:
                    if len(row) >= 7 and row[0] == str(num_boucle) and row[1] == dossard:
                        continue 
                    lignes.append(row)
                    
            with open(FICHIER_RESULTATS, mode="w", newline="", encoding="utf-8") as f:
                writer = csv.writer(f)
                writer.writerows(lignes)

            etat_course["logs_recents"].insert(0, {"text": f"Annulation : Dossard {dossard} retiré de la boucle {num_boucle}.", "color": "#f38ba8"})
            return jsonify({"status": "ok"})
        else:
            return jsonify({"status": "error", "message": f"Le dossard {dossard} n'a pas encore validé cette boucle."})

@app.route("/api/export_excel", methods=["GET"])
def export_excel():
    try:
        import openpyxl
        from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
        from openpyxl.drawing.image import Image
    except ImportError:
        return "Erreur : La librairie openpyxl ou Pillow n'est pas installée. Sur votre Mac, ouvrez le Terminal et tapez : pip install openpyxl Pillow", 500

    wb = openpyxl.Workbook()
    
    border_thin = Border(left=Side(style='thin', color='D3D3D3'), right=Side(style='thin', color='D3D3D3'), top=Side(style='thin', color='D3D3D3'), bottom=Side(style='thin', color='D3D3D3'))
    font_bold = Font(bold=True)
    font_header = Font(bold=True, color="FFFFFF", size=12)
    fill_header_main = PatternFill(start_color="313244", end_color="313244", fill_type="solid")
    fill_header_blue = PatternFill(start_color="89b4fa", end_color="89b4fa", fill_type="solid")
    fill_label = PatternFill(start_color="F2F2F7", end_color="F2F2F7", fill_type="solid")
    fill_best_lap = PatternFill(start_color="E8F8F5", end_color="E8F8F5", fill_type="solid")
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")

    ws_general = wb.active
    ws_general.title = "Classement Général"
    ws_general.sheet_view.showGridLines = False 

    chemin_logo_by = os.path.join("static", "logo_backyard.png")
    chemin_logo_tdr = os.path.join("static", "logo_tdr.png")
    chemin_logo_fp = os.path.join("static", "logo_fitness_park.png")
    
    for _ in range(5):
        ws_general.append([])
        
    ws_general.row_dimensions[1].height = 6
    for row_idx in range(2, 6):
        ws_general.row_dimensions[row_idx].height = 24
        
    if os.path.exists(chemin_logo_by):
        img_by = Image(chemin_logo_by)
        img_by.height = 90
        img_by.width = 90
        ws_general.add_image(img_by, 'B2') 
        
    if os.path.exists(chemin_logo_tdr):
        img_tdr = Image(chemin_logo_tdr)
        img_tdr.height = 90
        img_tdr.width = 90
        ws_general.add_image(img_tdr, 'G2') 

    if os.path.exists(chemin_logo_fp):
        img_fp = Image(chemin_logo_fp)
        img_fp.height = 90
        img_fp.width = 90
        ws_general.add_image(img_fp, 'H2') 

    ws_general.merge_cells('C1:F5')
    titre_main = ws_general['C1']
    titre_main.value = "RÉSULTATS OFFICIELS\nBACKYARD ULTRA PHELMA"
    titre_main.font = Font(bold=True, size=24, color="313244")
    titre_main.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    headers = ["Rang", "Dossard", "Nom", "Boucles validées", "Distance Totale (km)", "Meilleur Tour", "Moyenne par tour", "Allure Moyenne"]
    ws_general.append(headers)
    
    for col in range(1, 9):
        cell = ws_general.cell(row=6, column=col)
        cell.font = font_header
        cell.fill = fill_header_main
        cell.alignment = align_center
        cell.border = border_thin

    tours_tous = {}
    for d, nom in etat_course["coureurs"].items():
        tours_tous[d] = {"dossard": d, "nom": nom, "tours": 0, "total_sec": 0, "total_repos": 0, "fastest_sec": 999999, "fastest_str": "--:--", "last_loop_sec": 999999, "details": []}
        
    if os.path.exists(FICHIER_RESULTATS):
        with open(FICHIER_RESULTATS, mode="r", encoding="utf-8") as f:
            reader = csv.reader(f)
            next(reader, None)
            for row in reader:
                if len(row) >= 7:
                    b, d, nom, h_arr, t_boucle, t_repos, statut = row
                    if d in tours_tous:
                        tours_tous[d]["details"].append([b, t_boucle, t_repos, statut])
                        if "OK" in statut:
                            tours_tous[d]["tours"] += 1
                            try:
                                m, s = map(int, t_boucle.split(":"))
                                sec = m * 60 + s
                                tours_tous[d]["total_sec"] += sec
                                tours_tous[d]["last_loop_sec"] = sec
                                
                                repos_m, repos_s = map(int, t_repos.split(":"))
                                tours_tous[d]["total_repos"] += (repos_m * 60 + repos_s)
                                
                                if sec < tours_tous[d]["fastest_sec"]:
                                    tours_tous[d]["fastest_sec"] = sec
                                    tours_tous[d]["fastest_str"] = f"{t_boucle} (B{b})"
                            except:
                                pass

    def sort_excel(x):
        avg = (x["total_sec"] / x["tours"]) if x["tours"] > 0 else float('inf')
        return (-x["tours"], avg)
    
    classement = sorted(tours_tous.values(), key=sort_excel)
    
    row_num_gen = 7 
    for i, c in enumerate(classement):
        c["rang_final"] = i + 1 
        
        moy_str = "--:--"
        allure_str = "--:--"
        if c["tours"] > 0:
            avg = c["total_sec"] // c["tours"]
            moy_str = f"{avg//60:02d}:{avg%60:02d}"
            
            allure_sec = int(avg / DISTANCE_BOUCLE_KM)
            allure_str = f"{allure_sec//60:02d}:{allure_sec%60:02d}/km"

        c["moy_str_final"] = moy_str
        c["allure_str_final"] = allure_str
        
        try:
            dossard_num = int(c["dossard"])
        except ValueError:
            dossard_num = c["dossard"]
            
        row_data = [
            c["rang_final"], dossard_num, c["nom"], c["tours"],
            round(c["tours"] * DISTANCE_BOUCLE_KM, 1), c["fastest_str"], c["moy_str_final"], c["allure_str_final"]
        ]
        ws_general.append(row_data)
        
        for col_idx in range(1, 9):
            cell = ws_general.cell(row=row_num_gen, column=col_idx)
            cell.border = border_thin
            cell.alignment = align_center if col_idx != 3 else align_left
            
        row_num_gen += 1
        
    def get_dossard_sort_key(c_dict):
        try:
            return int(c_dict["dossard"])
        except ValueError:
            return 999999

    coureurs_tries_dossard = sorted(tours_tous.values(), key=get_dossard_sort_key)
    
    for c in coureurs_tries_dossard:
        if c["tours"] > 0 or len(c["details"]) > 0:
            safe_name = f"{c['dossard']} - {c['nom']}"[:31].replace(":", "").replace("/", "")
            ws_indiv = wb.create_sheet(title=safe_name)
            ws_indiv.sheet_view.showGridLines = False 
            
            ws_indiv.merge_cells('A1:E1')
            title_cell = ws_indiv['A1']
            title_cell.value = f"RÉSULTATS : {c['nom']} (Dossard {c['dossard']})"
            title_cell.font = Font(bold=True, size=16, color="FFFFFF")
            title_cell.fill = fill_header_main
            title_cell.alignment = align_center
            ws_indiv.row_dimensions[1].height = 30
            
            rh, r_rem = divmod(c["total_repos"], 3600)
            rm, _ = divmod(r_rem, 60)
            repos_str = f"{rh}h {rm:02d}m" if rh > 0 else f"{rm}m"
            
            stats_labels = [
                ("Classement Général", f"{c.get('rang_final', '-')} / {len(tours_tous)}"),
                ("Boucles validées", c["tours"]),
                ("Distance Officielle", f"{c['tours'] * DISTANCE_BOUCLE_KM:.1f} km"),
                ("Moyenne par tour", c.get("moy_str_final", "--:--")),
                ("Allure Moyenne Globale", c.get("allure_str_final", "--:--")),
                ("Meilleur tour", c["fastest_str"]),
                ("Temps de Repos Cumulé", repos_str)
            ]
            
            row_idx = 3
            for label, val in stats_labels:
                c_lbl = ws_indiv.cell(row=row_idx, column=1, value=label)
                c_lbl.font = font_bold
                c_lbl.fill = fill_label
                c_lbl.border = border_thin
                
                c_val = ws_indiv.cell(row=row_idx, column=2, value=val)
                c_val.alignment = align_left
                c_val.border = border_thin
                
                ws_indiv.merge_cells(start_row=row_idx, start_column=2, end_row=row_idx, end_column=5)
                ws_indiv.cell(row=row_idx, column=5).border = border_thin 
                row_idx += 1
            
            row_idx += 2
            
            headers_indiv = ["Boucle N°", "Temps du Tour", "Allure (min/km)", "Temps de Repos", "Statut"]
            for col_idx, h in enumerate(headers_indiv, 1):
                cell = ws_indiv.cell(row=row_idx, column=col_idx, value=h)
                cell.font = font_header
                cell.fill = fill_header_blue
                cell.alignment = align_center
                cell.border = border_thin
                
            row_idx += 1
            
            for d in c["details"]:
                b_val, t_val, t_repos_val, statut_val = d
                
                is_best = False
                allure_tour_str = "--:--"
                try:
                    m, s = map(int, t_val.split(":"))
                    total_sec_tour = m * 60 + s
                    
                    if total_sec_tour == c["fastest_sec"] and "OK" in statut_val:
                        is_best = True
                        
                    allure_sec = int(total_sec_tour / DISTANCE_BOUCLE_KM)
                    allure_tour_str = f"{allure_sec//60:02d}:{allure_sec%60:02d}"
                except:
                    pass
                
                row_to_write = [b_val, t_val, allure_tour_str, t_repos_val, statut_val]
                
                for col_idx, val in enumerate(row_to_write, 1):
                    cell = ws_indiv.cell(row=row_idx, column=col_idx, value=val)
                    cell.alignment = align_center
                    cell.border = border_thin
                    
                    if is_best:
                        cell.fill = fill_best_lap
                        if col_idx == 2 or col_idx == 3: 
                            cell.font = Font(color="229954", bold=True)
                            
                    if col_idx == 5: 
                        if "OK" in val:
                            cell.font = Font(color="008000", bold=True)
                        else:
                            cell.font = Font(color="FF0000", bold=True)
                            
                row_idx += 1
                
            ws_indiv.column_dimensions['A'].width = 28
            ws_indiv.column_dimensions['B'].width = 20
            ws_indiv.column_dimensions['C'].width = 20
            ws_indiv.column_dimensions['D'].width = 20
            ws_indiv.column_dimensions['E'].width = 20

    ws_general.column_dimensions['A'].width = 10
    ws_general.column_dimensions['B'].width = 15
    ws_general.column_dimensions['C'].width = 30
    ws_general.column_dimensions['D'].width = 20
    ws_general.column_dimensions['E'].width = 25
    ws_general.column_dimensions['F'].width = 20
    ws_general.column_dimensions['G'].width = 20
    ws_general.column_dimensions['H'].width = 20

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    
    timestamp = datetime.now().strftime("%Y%m%d_%H%M")
    filename = f"Backyard_Resultats_{timestamp}.xlsx"
    
    return send_file(out, download_name=filename, as_attachment=True, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

@app.route("/api/download_statuts", methods=["GET"])
def download_statuts():
    if os.path.exists("Suivi_Statuts_Coureurs.xlsx"):
        return send_file("Suivi_Statuts_Coureurs.xlsx", as_attachment=True)
    return "La première boucle n'est pas encore terminée, le fichier de statuts n'a pas encore été généré.", 404

@app.route("/api/coureur/<dossard>", methods=["GET"])
def get_coureur_stats(dossard):
    dossard = dossard.strip()
    stats_tours = []
    nom_coureur = etat_course["coureurs"].get(dossard, "Inconnu")
    
    tours_tous = {d: {"tours": 0, "last_loop_sec": 999999, "total_sec": 0} for d in etat_course["coureurs"]}
    
    total_sec = 0
    total_repos_sec = 0
    fastest_sec = 999999
    fastest_str = "--:--"
    fastest_boucle = "-" 
    tours_valides = 0

    if os.path.exists(FICHIER_RESULTATS):
        with open(FICHIER_RESULTATS, mode="r", encoding="utf-8") as f:
            reader = csv.reader(f)
            next(reader, None) 
            for row in reader:
                if len(row) >= 7:
                    b = row[0]
                    d = row[1]
                    t = row[4]
                    statut = row[6]

                    if "OK" in statut and d in tours_tous:
                        tours_tous[d]["tours"] += 1
                        try:
                            m, s = map(int, t.split(":"))
                            sec = m * 60 + s
                            tours_tous[d]["last_loop_sec"] = sec
                            tours_tous[d]["total_sec"] += sec
                        except:
                            pass

                    if d == dossard:
                        stats_tours.append({
                            "boucle": b, "temps": t, "statut": statut
                        })
                        if "OK" in statut:
                            tours_valides += 1
                            try:
                                m, s = map(int, t.split(":"))
                                t_sec = m * 60 + s
                                total_sec += t_sec
                                repos = 3600 - t_sec if t_sec < 3600 else 0
                                total_repos_sec += repos

                                if t_sec < fastest_sec:
                                    fastest_sec = t_sec
                                    fastest_str = t
                                    fastest_boucle = b 
                            except:
                                pass

    if not stats_tours and dossard not in etat_course["coureurs"]:
        return jsonify({"error": f"Aucun dossard trouvé avec le numéro {dossard}."})

    def sort_class(x):
        avg = (x[1]["total_sec"] / x[1]["tours"]) if x[1]["tours"] > 0 else float('inf')
        return (-x[1]["tours"], avg)

    classement_trie = sorted(tours_tous.items(), key=sort_class)
    rang = 0
    for i, (d_id, data) in enumerate(classement_trie):
        if d_id == dossard:
            rang = i + 1
            break
    total_coureurs = len(tours_tous)

    if tours_valides > 0:
        moy_sec = total_sec // tours_valides
        moy_str = f"{moy_sec//60:02d}:{moy_sec%60:02d}"
        allure_sec = int(moy_sec / DISTANCE_BOUCLE_KM)
        allure_str = f"{allure_sec//60:02d}:{allure_sec%60:02d}/km"
    else:
        moy_str = "--:--"
        allure_str = "--:--"
        
    if fastest_str != "--:--":
        fastest_display = f"{fastest_str} (B{fastest_boucle})"
    else:
        fastest_display = "--:--"

    rh, r_rem = divmod(total_repos_sec, 3600)
    rm, _ = divmod(r_rem, 60)
    repos_str = f"{rh}h {rm:02d}m" if rh > 0 else f"{rm}m"

    distance_str = f"{tours_valides * DISTANCE_BOUCLE_KM:.1f} km"

    return jsonify({
        "nom": nom_coureur,
        "dossard": dossard,
        "rang": rang,
        "total_coureurs": total_coureurs,
        "tours_total": tours_valides,
        "distance": distance_str,
        "repos_total": repos_str,
        "fastest": fastest_display, 
        "fastest_boucle_num": str(fastest_boucle),
        "moyenne": moy_str,
        "allure": allure_str,
        "details": stats_tours
    })

@app.route("/api/classement", methods=["GET"])
def get_classement():
    tours_tous = {}
    for d, nom in etat_course["coureurs"].items():
        tours_tous[d] = {"dossard": d, "nom": nom, "tours": 0, "total_sec": 0, "fastest_sec": 999999, "fastest_str": "--:--", "last_loop_sec": 999999}
        
    if os.path.exists(FICHIER_RESULTATS):
        with open(FICHIER_RESULTATS, mode="r", encoding="utf-8") as f:
            reader = csv.reader(f)
            next(reader, None)
            for row in reader:
                if len(row) >= 7:
                    b = row[0]
                    d = row[1]
                    t = row[4]
                    statut = row[6]
                    if "OK" in statut and d in tours_tous:
                        tours_tous[d]["tours"] += 1
                        try:
                            m, s = map(int, t.split(":"))
                            sec = m * 60 + s
                            tours_tous[d]["total_sec"] += sec
                            tours_tous[d]["last_loop_sec"] = sec
                            if sec < tours_tous[d]["fastest_sec"]:
                                tours_tous[d]["fastest_sec"] = sec
                                tours_tous[d]["fastest_str"] = f"{t} (B{b})"
                        except:
                            pass
                            
    def sort_class(x):
        avg = (x["total_sec"] / x["tours"]) if x["tours"] > 0 else float('inf')
        return (-x["tours"], avg)

    classement = sorted(tours_tous.values(), key=sort_class)
    
    resultat = []
    for i, c in enumerate(classement):
        moy_str = "--:--"
        if c["tours"] > 0:
            avg = c["total_sec"] // c["tours"]
            moy_str = f"{avg//60:02d}:{avg%60:02d}"
            
        resultat.append({
            "rang": i + 1,
            "dossard": c["dossard"],
            "nom": c["nom"],
            "tours": c["tours"],
            "distance": f"{c['tours'] * DISTANCE_BOUCLE_KM:.1f} km",
            "fastest": c["fastest_str"],
            "moyenne": moy_str
        })
        
    return jsonify({"classement": resultat})

@app.route("/api/podium", methods=["GET"])
def get_podium():
    tours_valides = {}
    fastest_sec = 999999
    fastest_runner = None
    fastest_str = "--:--"
    
    if not os.path.exists(FICHIER_RESULTATS):
        return jsonify({"error": "Aucune donnée"})

    with open(FICHIER_RESULTATS, mode="r", encoding="utf-8") as f:
        reader = csv.reader(f)
        next(reader, None) 
        for row in reader:
            if len(row) >= 7:
                b = int(row[0])
                d = row[1]
                nom = row[2]
                t = row[4]
                statut = row[6]
                
                if d not in tours_valides:
                    tours_valides[d] = {"dossard": d, "nom": nom, "tours": 0, "total_sec": 0, "last_loop_sec": 999999, "last_loop_str": "--:--"}
                
                if "OK" in statut:
                    tours_valides[d]["tours"] += 1
                    try:
                        m, s = map(int, t.split(":"))
                        sec = m * 60 + s
                        tours_valides[d]["total_sec"] += sec
                        tours_valides[d]["last_loop_sec"] = sec
                        tours_valides[d]["last_loop_str"] = t
                        
                        if sec < fastest_sec:
                            fastest_sec = sec
                            fastest_str = f"{t} (Boucle {b})"
                            fastest_runner = f"#{d} - {nom}"
                    except:
                        pass

    def sort_podium(x):
        avg = (x["total_sec"] / x["tours"]) if x["tours"] > 0 else float('inf')
        return (-x["tours"], avg)

    classement = sorted(tours_valides.values(), key=sort_podium)
    
    vainqueurs = []
    assists = []
    is_tie = False
    
    if classement and classement[0]["tours"] > 0:
        max_tours = classement[0]["tours"]
        candidats = [c for c in classement if c["tours"] == max_tours]
        
        if len(candidats) > 1:
            is_tie = True
            candidats = sorted(candidats, key=lambda x: (x["total_sec"]/x["tours"]) if x["tours"] > 0 else float('inf'))
            vainqueurs = [candidats[0]] 
            assists = [candidats[1]]    
        else:
            vainqueurs = candidats
            assists_candidates = [c for c in classement if c["tours"] < max_tours]
            if assists_candidates:
                max_assist = assists_candidates[0]["tours"]
                assists_ex_aequo = [c for c in assists_candidates if c["tours"] == max_assist]
                if len(assists_ex_aequo) > 1:
                    assists_ex_aequo = sorted(assists_ex_aequo, key=lambda x: (x["total_sec"]/x["tours"]) if x["tours"] > 0 else float('inf'))
                assists = [assists_ex_aequo[0]]

    best_avg_sec = 999999
    best_avg_runner = None
    best_avg_str = "--:--"
    
    for c in classement:
        if c["tours"] > 0:
            avg = c["total_sec"] // c["tours"]
            if avg < best_avg_sec:
                best_avg_sec = avg
                best_avg_runner = f"#{c['dossard']} - {c['nom']}"
                best_avg_str = f"{avg//60:02d}:{avg%60:02d}"

    def formater_liste(liste, montrer_sprint=False):
        resultat = []
        for r in liste:
            detail = f"{r['tours']} boucles ({r['tours']*DISTANCE_BOUCLE_KM:.1f} km)"
            if montrer_sprint:
                detail += f"<br><span style='font-size:0.8em; color:#f9e2af;'>Dernier tour : {r['last_loop_str']}</span>"
            resultat.append({"texte": f"#{r['dossard']} - {r['nom']}", "detail": detail})
        return resultat

    return jsonify({
        "vainqueurs": formater_liste(vainqueurs, is_tie),
        "assists": formater_liste(assists, is_tie),
        "record_tour": {"runner": fastest_runner, "detail": fastest_str} if fastest_runner else None,
        "meilleure_moy": {"runner": best_avg_runner, "detail": best_avg_str} if best_avg_runner else None
    })

# ==========================================
# PAGES WEB (HTML)
# ==========================================

HTML_MAIN = """
<!DOCTYPE html>
<html lang="fr">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Chrono Backyard Ultra</title>
    <style>
        body { font-family: 'Helvetica', Arial, sans-serif; background-color: #1e1e2e; color: #cdd6f4; text-align: center; margin: 0; padding: 15px 30px; min-height: 100vh; box-sizing: border-box; display: flex; flex-direction: column; overflow-y: auto; }
        .top-header { display: flex; justify-content: space-between; align-items: center; margin-bottom: 20px; width: 100%; }
        .header-left { width: 300px; text-align: left; }
        .header-center { flex: 1; text-align: center; }
        .header-right { width: 300px; display: flex; gap: 15px; justify-content: flex-end; align-items: center; }
        .top-header img { height: 100px; object-fit: contain; border-radius: 8px; }
        h1 { margin: 5px 0; font-size: 2.2em; }
        .clock { font-size: 3.5em; font-weight: bold; margin: 10px 0; color: #f9e2af; }
        .danger { color: #f38ba8 !important; } .safe { color: #a6e3a1 !important; }
        input[type="text"], input[type="password"] { font-size: 2em; padding: 10px; width: 150px; text-align: center; font-weight: bold; border: 2px solid #313244; border-radius: 8px; background: #11111b; color: white; margin-bottom: 10px;}
        button { font-size: 1.1em; padding: 10px 20px; margin: 5px; cursor: pointer; border: none; border-radius: 8px; font-weight: bold; }
        .btn-start { background-color: #a6e3a1; color: #11111b; font-size: 1.5em; padding: 15px 30px; margin-top: 50px; }
        .btn-correction { background-color: #f38ba8; color: #11111b; }
        .btn-undo { background-color: #fab387; color: #11111b; }
        #inline-msg { height: 25px; font-size: 1.2em; font-weight: bold; margin-bottom: 10px; }
        .logs { background-color: #181825; padding: 15px; border-radius: 8px; max-width: 650px; width: 100%; margin: 15px auto 0 auto; text-align: left; font-family: monospace; font-size: 1.1em; flex: 1; min-height: 200px; box-sizing: border-box; }
        .hidden { display: none !important; }
        #lock-screen { display: flex; flex-direction: column; justify-content: center; align-items: center; flex: 1; }
        #lock-screen h2 { color: #f38ba8; font-size: 2em;}
        #btn-unlock { background-color: #89b4fa; color: #11111b; font-size: 1.5em; padding: 10px 30px;}
        @media (max-width: 768px) {
            body { padding: 10px; }
            .top-header { flex-direction: column; gap: 10px; }
            .top-header img { height: 70px; }
            h1 { font-size: 1.8em; }
            .clock { font-size: 2.8em; }
            .logs { min-height: 400px; font-size: 1em; }
        }
    </style>
</head>
<body>
    <div class="top-header">
        <div class="header-left">
            <img src="/static/logo_backyard.png" alt="Logo Phoenix">
        </div>
        <div class="header-center">
            <h1>Chronométrage Backyard Ultra</h1>
            <a href="/roulette" target="_blank" style="display: inline-block; background: #89b4fa; color: #11111b; padding: 8px 15px; border-radius: 8px; text-decoration: none; font-weight: bold; margin-top: 5px;">Ouvrir la Roulette</a>
        </div>
        <div class="header-right">
            <img src="/static/logo_tdr.png" alt="Logo TDR">
            <img src="/static/logo_fitness_park.png" alt="Logo Fitness Park">
        </div>
    </div>
    
    <div id="lock-screen">
        <h2>Verrouillage de sécurité</h2>
        <p>Veuillez entrer le code PIN pour accéder au pupitre de chronométrage.</p>
        <input type="password" id="pin-input" placeholder="PIN" autocomplete="off" maxlength="4">
        <button id="btn-unlock" onclick="checkPin()">Déverrouiller</button>
        <div id="pin-error" style="color: #f38ba8; margin-top: 10px; font-weight: bold;"></div>
    </div>
    
    <div id="main-app" class="hidden" style="flex: 1; display: flex; flex-direction: column;">
        <div id="start-section" style="margin-top: 30px;">
            <button class="btn-start" onclick="startRace()">LANCER LA COURSE</button>
            <div style="margin-top: 60px; background: #181825; padding: 20px; border-radius: 12px; display: inline-block;">
                <p style="color: #a6adc8; margin-top: 0; font-weight: bold;">Ou reprendre une course depuis un backup (.csv) :</p>
                <input type="file" id="backup-file" accept=".csv" style="font-size: 0.9em; margin-bottom: 15px; color: white;">
                <br>
                <button class="btn-undo" onclick="restoreBackup()" style="font-size: 1em;">Restaurer le fichier</button>
            </div>
        </div>
        
        <div id="race-section" class="hidden">
            <h2 id="loop-info" style="margin: 5px 0;">Boucle -- en cours</h2>
            <div class="clock" id="countdown">--:--</div>
            
            <div id="inline-msg"></div>

            <form id="chrono-form" onsubmit="submitDossard(event, false)" style="margin-bottom: 5px;">
                <label style="font-size: 1.2em; display: block; margin-bottom: 8px;">Dossard :</label>
                <input type="text" id="dossard-input" autocomplete="off">
            </form>
            
            <div>
                <button class="btn-correction" onclick="promptCorrection()">Oubli : Repêcher</button>
                <button class="btn-undo" onclick="promptUndo()">Annuler dossard</button>
            </div>
        </div>
        
        <div class="logs" id="logs-container"></div>
    </div>
    
<script>
    const dossardInput = document.getElementById("dossard-input");
    const pinInput = document.getElementById("pin-input");
    let isUnlocked = false;

    pinInput.addEventListener("keypress", function(event) {
        if (event.key === "Enter") { event.preventDefault(); checkPin(); }
    });

    async function checkPin() {
        const pin = pinInput.value;
        const res = await fetch("/api/verify_pin", {
            method: "POST", headers: {"Content-Type": "application/json"},
            body: JSON.stringify({pin: pin})
        });
        const result = await res.json();
        
        if (result.status === "ok") {
            isUnlocked = true;
            document.getElementById("lock-screen").classList.add("hidden");
            document.getElementById("main-app").classList.remove("hidden");
            fetchState(); 
            if(audioCtx.state === 'suspended') audioCtx.resume();
        } else {
            document.getElementById("pin-error").innerText = result.message;
            pinInput.value = "";
        }
    }

    const audioCtx = new (window.AudioContext || window.webkitAudioContext)();

    function playBeep(type) {
        if(audioCtx.state === 'suspended') audioCtx.resume();
        const osc = audioCtx.createOscillator();
        const gain = audioCtx.createGain();
        osc.connect(gain);
        gain.connect(audioCtx.destination);

        if(type === 'success') {
            osc.type = 'sine'; 
            osc.frequency.setValueAtTime(880, audioCtx.currentTime); 
            gain.gain.setValueAtTime(0.5, audioCtx.currentTime);
            osc.start();
            gain.gain.exponentialRampToValueAtTime(0.01, audioCtx.currentTime + 0.15);
            osc.stop(audioCtx.currentTime + 0.15);
        } else {
            osc.type = 'sawtooth'; 
            osc.frequency.setValueAtTime(150, audioCtx.currentTime);
            gain.gain.setValueAtTime(0.5, audioCtx.currentTime);
            osc.start();
            gain.gain.exponentialRampToValueAtTime(0.01, audioCtx.currentTime + 0.4);
            osc.stop(audioCtx.currentTime + 0.4);
        }
    }
    
    async function fetchState(){
        if (!isUnlocked) return; 
        
        const res=await fetch("/api/state");
        const data=await res.json();
        if(data.started){
            document.getElementById("start-section").classList.add("hidden");
            document.getElementById("race-section").classList.remove("hidden");
            document.getElementById("loop-info").innerText=`BOUCLE ${data.boucle} (${data.survivants_total} en lice)`;
            let sec=data.temps_restant;
            let clockEl=document.getElementById("countdown");
            if(sec>0){
                let m=Math.floor(sec/60).toString().padStart(2,'0');
                let s=(sec%60).toString().padStart(2,'0');
                clockEl.innerText=`Départ dans : ${m}:${s}`;
                clockEl.className=sec<180?"clock danger":"clock";
            } else {
                clockEl.innerText="TOP DÉPART !";
                clockEl.className="clock safe";
            }
        }
        document.getElementById("logs-container").innerHTML=data.logs.map(log=>`<div style="color: ${log.color}; margin-bottom: 5px;">${log.text}</div>`).join('');
    }

    async function startRace(){
        if(confirm("Démarrer la course ?")){
            if(audioCtx.state === 'suspended') audioCtx.resume();
            await fetch("/api/start",{method:"POST"});
            fetchState();
        }
    }

    async function restoreBackup() {
        const fileInput = document.getElementById('backup-file');
        if (!fileInput.files.length) {
            alert("Veuillez sélectionner un fichier CSV sur votre ordinateur.");
            return;
        }
        if (confirm("Attention, l'envoi de ce fichier va écraser la course en cours de mémoire. Continuer ?")) {
            const formData = new FormData();
            formData.append("file", fileInput.files[0]);
            
            const res = await fetch("/api/restore", {
                method: "POST",
                body: formData
            });
            const result = await res.json();
            if (result.status === "ok") {
                alert("Restauration réussie ! La course a été rechargée.");
                fileInput.value = "";
                fetchState();
            } else {
                alert("Erreur : " + result.message);
            }
        }
    }

    async function submitDossard(event, isCorrection, overrideDossard=null){
        if(event) event.preventDefault();
        let num=overrideDossard || dossardInput.value;
        if(!num) return;
        
        const res=await fetch("/api/submit",{method:"POST",headers:{"Content-Type":"application/json"},body:JSON.stringify({dossard:num,correction:isCorrection})});
        const result=await res.json();
        
        if (result.status !== "ok") {
            playBeep("error");
            const msgEl = document.getElementById("inline-msg");
            msgEl.innerText = result.message;
            msgEl.style.color = result.status === "duplicate" ? "#fab387" : "#f38ba8";
            setTimeout(() => msgEl.innerText = "", 3000); 
        } else {
            playBeep("success");
        }
        
        dossardInput.value="";
        dossardInput.focus();
        fetchState();
    }

    function promptCorrection(){
        let num=prompt("Numéro du dossard à repêcher :");
        if(num) submitDossard(null,true,num);
        dossardInput.focus();
    }
    
    function promptUndo(){
        let num = prompt("Numéro du dossard à annuler pour la boucle actuelle :");
        if(num){
            fetch("/api/undo", {method: "POST", headers: {"Content-Type": "application/json"}, body: JSON.stringify({dossard: num})})
            .then(r => r.json()).then(res => {
                if(res.status !== "ok") {
                    playBeep("error");
                    const msgEl = document.getElementById("inline-msg");
                    msgEl.innerText = res.message;
                    msgEl.style.color = "#f38ba8";
                    setTimeout(() => msgEl.innerText = "", 3000);
                } else {
                    playBeep("success");
                    fetchState();
                }
            });
        }
        dossardInput.focus();
    }

    setInterval(fetchState,1000);
</script></body></html>
"""

HTML_ROULETTE = """
<!DOCTYPE html>
<html lang="fr">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Roulette des Lots - Backyard Ultra</title>
    <style>
        body { font-family: 'Helvetica', Arial, sans-serif; background-color: #1e1e2e; color: #cdd6f4; text-align: center; margin: 0; padding: 30px; display: flex; flex-direction: column; align-items: center; min-height: 100vh;}
        .header-logos { display: flex; justify-content: center; gap: 40px; margin-bottom: 20px; align-items: center; width: 100%; max-width: 1200px; }
        .header-logos img { height: 100px; border-radius: 12px;}
        h1 { font-size: 3em; color: #f9e2af; margin-top: 0; margin-bottom: 30px; text-transform: uppercase; letter-spacing: 2px; }
        .roulette-container { background: #181825; padding: 40px; border-radius: 20px; box-shadow: 0 10px 30px rgba(0,0,0,0.5); width: 100%; max-width: 600px; border: 4px solid #313244; z-index: 10;}
        .winner-display { font-size: 3em; font-weight: bold; height: 1.5em; line-height: 1.5em; background: #11111b; border-radius: 12px; margin-bottom: 30px; color: #ffffff; overflow: hidden; position: relative;}
        .btn-spin { font-size: 1.5em; padding: 15px 40px; cursor: pointer; background: #a6e3a1; color: #11111b; border: none; border-radius: 12px; font-weight: bold; text-transform: uppercase; transition: transform 0.1s;}
        .btn-spin:active { transform: scale(0.95); }
        .btn-spin:disabled { background: #585b70; cursor: not-allowed; transform: none;}
        .confetti { position: absolute; width: 10px; height: 10px; z-index: 1000; }
        #winner-name { display: inline-block; width: 100%; transition: transform 0.1s ease-out; }
        
        .winners-history { margin-top: 30px; background: #181825; padding: 20px; border-radius: 16px; width: 100%; max-width: 1200px; box-sizing: border-box;}
        .winners-history h2 { color: #89b4fa; margin-top: 0; font-size: 1.8em; border-bottom: 2px solid #313244; padding-bottom: 10px;}
        #winners-list { list-style: none; padding: 0; margin: 0; font-size: 1.3em; font-weight: bold; color: #a6e3a1; display: flex; flex-direction: row; flex-wrap: wrap; justify-content: center; gap: 15px; }
        #winners-list li { background: #313244; padding: 10px 20px; border-radius: 8px; animation: fadeIn 0.5s ease-out;}
        @keyframes fadeIn { from { opacity: 0; transform: translateY(-10px); } to { opacity: 1; transform: translateY(0); } }
    </style>
</head>
<body>
    <div class="header-logos">
        <img src="/static/logo_backyard.png" alt="Logo">
        <img src="/static/logo_tdr.png" alt="Logo TDR">
        <img src="/static/logo_fitness_park.png" alt="Logo Fitness Park">
    </div>
    
    <h1>Tirage au sort</h1>
    
    <div class="roulette-container">
        <div class="winner-display">
            <span id="winner-name">Prêt à tirer !</span>
        </div>
        <button class="btn-spin" id="btn-spin" onclick="spin()">Lancer la roulette</button>
    </div>
    
    <div class="winners-history">
        <h2>Historique des gagnants</h2>
        <ul id="winners-list">
            <li style="color: #a6adc8; font-weight: normal; background: transparent; font-size: 0.8em; font-style: italic; width: 100%;" id="empty-state">Aucun gagnant pour le moment</li>
        </ul>
    </div>
    
    <script>
        let participants = [];
        
        fetch('/api/participants').then(r => r.json()).then(data => {
            participants = data.participants;
        });

        function spin() {
            if (participants.length === 0) return alert("Il n'y a plus de participants éligibles pour le tirage !");
            
            const btn = document.getElementById("btn-spin");
            const display = document.getElementById("winner-name");
            
            btn.disabled = true;
            display.style.color = "#ffffff";
            display.style.transform = "scale(1)";

            let delay = 50;
            let totalSpins = 40; 
            let currentSpin = 0;
            let winningIndex = 0;

            function tick() {
                winningIndex = Math.floor(Math.random() * participants.length);
                let p = participants[winningIndex];
                display.innerText = `#${p.dossard} - ${p.nom}`;
                
                currentSpin++;
                if (currentSpin < totalSpins) {
                    if (currentSpin > 25) delay += 15;
                    setTimeout(tick, delay);
                } else {
                    btn.disabled = false;
                    display.style.color = "#a6e3a1";
                    display.style.transform = "scale(1.1)";
                    setTimeout(() => {
                        display.style.transform = "scale(1)";
                        ajouterGagnantALaListe(display.innerText);
                    }, 300);
                    shootConfetti();
                    
                    // Retrait immédiat du gagnant de la liste
                    participants.splice(winningIndex, 1);
                }
            }
            tick();
        }
        
        function ajouterGagnantALaListe(nom) {
            const ul = document.getElementById("winners-list");
            const emptyState = document.getElementById("empty-state");
            if (emptyState) emptyState.remove();
            
            const li = document.createElement("li");
            li.innerText = nom;
            ul.prepend(li);
        }

        function shootConfetti() {
            for(let i=0; i<60; i++) {
                let conf = document.createElement('div');
                conf.classList.add('confetti');
                conf.style.left = Math.random() * 100 + 'vw';
                conf.style.top = '-10px';
                conf.style.backgroundColor = ['#a6e3a1', '#f38ba8', '#89b4fa', '#f9e2af'][Math.floor(Math.random()*4)];
                document.body.appendChild(conf);
                
                let fallDuration = Math.random() * 2 + 2;
                conf.animate([
                    { transform: `translate3d(0,0,0) rotate(0deg)` },
                    { transform: `translate3d(${Math.random()*150 - 75}px, 100vh, 0) rotate(${Math.random()*720}deg)` }
                ], {
                    duration: fallDuration * 1000,
                    easing: 'linear'
                }).onfinish = () => conf.remove();
            }
        }
    </script>
</body>
</html>
"""

HTML_TV = """
<!DOCTYPE html>
<html lang="fr">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>TV Backyard - Leaderboard Dynamique</title>
    <style>
        body { font-family: 'Helvetica', Arial, sans-serif; background-color: #11111b; color: #cdd6f4; margin: 0; overflow: hidden; display: flex; flex-direction: column; height: 100vh; width: 100vw; }
        
        /* HEADER (12% hauteur) */
        .header { display: flex; justify-content: space-between; align-items: center; padding: 1vh 2vw; background: #1e1e2e; border-bottom: 0.3vh solid #313244; height: 12vh; box-sizing: border-box;}
        .header-left { width: 25vw; text-align: left; }
        .header-right { width: 25vw; display: flex; gap: 1vw; justify-content: flex-end; align-items: center; }
        .header-center { flex-grow: 1; text-align: center; }
        .header img { height: 8vh; object-fit: contain; border-radius: 1vh;}
        .title { font-size: 3.5vh; font-weight: bold; color: #ffffff; margin: 0 0 0.5vh 0; }
        .clock { font-size: 5vh; font-weight: bold; color: #f9e2af; margin: 0; line-height: 1;}
        .danger { color: #f38ba8 !important; }
        
        /* PODIUM BAR (5% hauteur) */
        .podium-bar { display: flex; justify-content: space-around; align-items: center; background-color: #89b4fa; color: #11111b; padding: 0 1vw; font-weight: bold; font-size: 1.8vh; height: 5vh; box-sizing: border-box; }
        
        /* TOOLS BAR (6% hauteur) */
        .tools-bar { display: flex; justify-content: space-between; align-items: center; padding: 0 1.5vw; height: 6vh; box-sizing: border-box;}
        .legende { display: flex; gap: 1vw; font-size: 1.6vh; font-weight: bold;}
        .leg-box { padding: 0.5vh 0.8vw; border-radius: 0.5vh; border-left: 0.6vh solid; background: #181825;}
        .page-indicator { font-size: 1.8vh; font-weight: bold; color: #f9e2af; background: #313244; padding: 0.5vh 1vw; border-radius: 1vh;}
        
        /* GRID (Reste de l'écran) */
        .grid-container { flex-grow: 1; padding: 0 1.5vw 1.5vh 1.5vw; overflow: hidden; display: flex; flex-direction: column; box-sizing: border-box;}
        .grid-all { 
            display: grid; 
            grid-template-columns: repeat(9, 1fr); 
            grid-template-rows: repeat(7, minmax(0, 1fr)); 
            gap: 0.8vh 0.5vw; 
            flex-grow: 1;
            height: 100%;
        }
        
        .card { background-color: #181825; padding: 0.8vh 0.5vw; border-radius: 0.8vh; border-left: 0.8vh solid; display: flex; flex-direction: column; justify-content: center; min-height: 0; overflow: hidden;}
        .card-head { font-weight: bold; font-size: 1.7vh; color: #ffffff; white-space: nowrap; overflow: hidden; text-overflow: ellipsis; margin-bottom: 0.5vh; }
        .card-stats { display: flex; justify-content: space-between; font-size: 1.2vh; color: #a6adc8; font-family: monospace; font-weight: bold;}
        
        .arrive { border-color: #a6e3a1; background-color: rgba(166, 227, 161, 0.15); } 
        .encourse { border-color: #fab387; } 
        .dnf { border-color: #f38ba8; opacity: 0.6; background-color: rgba(243, 139, 168, 0.05); } 
        
        @media (max-width: 768px) {
            body { overflow-y: auto !important; height: auto; display: block; }
            .grid-container { overflow: visible; height: auto; padding: 10px; }
            .grid-all { grid-template-columns: repeat(2, 1fr); grid-template-rows: auto; gap: 8px;}
            .header { height: auto; padding: 10px; flex-direction: column; gap: 10px; }
            .header-left, .header-right { width: 100%; justify-content: center; text-align: center; }
            .header img { height: 60px; }
            .clock { font-size: 3.5em; }
            .title { font-size: 1.5em; }
            .podium-bar { height: auto; flex-direction: column; font-size: 1.1em; text-align: center; gap: 5px; padding: 10px; }
            .tools-bar { height: auto; flex-direction: column; gap: 10px; padding: 10px;}
            .card { padding: 8px; border-left-width: 4px; }
            .card-head { font-size: 1.2em; margin-bottom: 4px;}
            .card-stats { font-size: 0.9em; flex-direction: column; text-align: center;}
            .page-indicator { display: none !important; }
        }
    </style>
</head>
<body>
    <div class="header">
        <div class="header-left">
            <img src="/static/logo_backyard.png" alt="Logo Phoenix">
        </div>
        <div class="header-center">
            <div class="title" id="loop-info">EN ATTENTE...</div>
            <div class="clock" id="countdown">--:--</div>
        </div>
        <div class="header-right">
            <img src="/static/logo_tdr.png" alt="Logo TDR">
            <img src="/static/logo_fitness_park.png" alt="Logo Fitness Park">
        </div>
    </div>
    
    <div class="podium-bar">
        <div>Record absolu : <span id="absolute-record">--</span></div>
        <div>Actuel plus rapide : <span id="fastest-loop">--</span></div>
        <div>Meilleure moy : <span id="best-avg">--</span></div>
    </div>

    <div class="tools-bar">
        <div class="legende">
            <div class="leg-box" style="border-color:#a6e3a1; background-color: rgba(166, 227, 161, 0.15);">Arrivé (Repos)</div>
            <div class="leg-box" style="border-color:#fab387;">En course</div>
            <div class="leg-box" style="border-color:#f38ba8; opacity: 0.6;">Éliminé</div>
        </div>
        <div class="page-indicator" id="page-info">Page 1/1</div>
    </div>
    
    <div class="grid-container">
        <div class="grid-all" id="list-all"></div>
    </div>

    <script>
        let globalRunners = [];
        let currentPage = 0;
        let CARDS_PER_PAGE = window.innerWidth <= 768 ? 10000 : 63; 

        window.addEventListener('resize', () => {
            let newPerPage = window.innerWidth <= 768 ? 10000 : 63;
            if(newPerPage !== CARDS_PER_PAGE) {
                CARDS_PER_PAGE = newPerPage;
                currentPage = 0;
                renderCurrentPage();
            }
        });

        async function fetchState() {
            const res = await fetch("/api/state");
            const data = await res.json();
            
            if (data.started) {
                document.getElementById("loop-info").innerText = `BOUCLE ${data.boucle} (${data.survivants_total} en lice)`;
                let sec = data.temps_restant;
                let clockEl = document.getElementById("countdown");
                if (sec > 0) {
                    let m = Math.floor(sec / 60).toString().padStart(2, '0');
                    let s = (sec % 60).toString().padStart(2, '0');
                    clockEl.innerText = `${m}:${s}`;
                    clockEl.className = sec < 180 ? "clock danger" : "clock";
                } else {
                    clockEl.innerText = "00:00";
                }

                globalRunners = data.runners_all;
                
                if (data.record_absolu && data.record_absolu.dossard) {
                    document.getElementById("absolute-record").innerText = `${data.record_absolu.str} (#${data.record_absolu.dossard} - ${data.record_absolu.nom})`;
                } else {
                    document.getElementById("absolute-record").innerText = `--`;
                }

                let arrivesThisLoop = globalRunners.filter(r => r.status === 'arrive');
                if(arrivesThisLoop.length > 0) {
                    let fastest = arrivesThisLoop.reduce((prev, curr) => prev.last_loop_sec < curr.last_loop_sec ? prev : curr);
                    document.getElementById("fastest-loop").innerText = `${fastest.last_loop} (#${fastest.dossard} - ${fastest.nom})`;
                } else {
                    document.getElementById("fastest-loop").innerText = `--`;
                }

                let validAvg = globalRunners.filter(r => r.moyenne_sec < 999999);
                if(validAvg.length > 0) {
                    let bestAvg = validAvg.reduce((prev, curr) => prev.moyenne_sec < curr.moyenne_sec ? prev : curr);
                    document.getElementById("best-avg").innerText = `${bestAvg.moyenne} (#${bestAvg.dossard} - ${bestAvg.nom})`;
                }
                renderCurrentPage();
            }
        }

        function renderCurrentPage() {
            if(globalRunners.length === 0) return;
            let totalPages = Math.ceil(globalRunners.length / CARDS_PER_PAGE);
            if (currentPage >= totalPages) currentPage = 0;
            document.getElementById("page-info").innerText = `Page ${currentPage + 1}/${totalPages}`;

            let startIdx = currentPage * CARDS_PER_PAGE;
            let endIdx = startIdx + CARDS_PER_PAGE;
            
            document.getElementById("list-all").innerHTML = globalRunners.slice(startIdx, endIdx).map(r => {
                if (r.status === 'dnf') {
                    let motBoucle = r.tours_valides > 1 ? "b" : "b"; 
                    return `
                    <div class="card ${r.status}">
                        <div class="card-head">#${r.dossard} - ${r.nom}</div>
                        <div class="card-stats" style="justify-content: center; color: #f38ba8; font-size: 1.1em; padding-top: 2px;">
                            Terminé : ${r.tours_valides}${motBoucle}
                        </div>
                    </div>`;
                } else {
                    return `
                    <div class="card ${r.status}">
                        <div class="card-head">#${r.dossard} - ${r.nom}</div>
                        <div class="card-stats">
                            <span title="Dernier Tour">${r.last_loop}</span>
                            <span title="Moyenne">Ø ${r.moyenne}</span>
                            <span title="Allure km">${r.allure}</span>
                        </div>
                    </div>`;
                }
            }).join('');
        }

        setInterval(fetchState, 1000);
        fetchState();
        setInterval(() => {
            if(globalRunners.length > 0 && window.innerWidth > 768) {
                currentPage = (currentPage + 1) % Math.ceil(globalRunners.length / CARDS_PER_PAGE);
                renderCurrentPage();
            }
        }, 10000);
    </script>
</body>
</html>
"""

HTML_COUREUR = """
<!DOCTYPE html>
<html lang="fr">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0, maximum-scale=1.0, user-scalable=no">
    <title>Borne Coureur</title>
    <style>
        body { font-family: 'Helvetica', sans-serif; background-color: #1e1e2e; color: #cdd6f4; margin: 0; padding: 0; height: 100vh; display: flex; flex-direction: column; overflow: hidden; }
        
        #view-search { flex: 1; display: flex; flex-direction: column; justify-content: center; align-items: center; padding: 20px; }
        .logos-container { display: flex; gap: 40px; margin-bottom: 30px; justify-content: center; align-items: center; }
        .logos-container img { height: 140px; border-radius: 12px; }
        h1 { font-size: 3em; margin-bottom: 30px; color: #ffffff; text-align: center;}
        
        input[type="number"] { font-size: 2.5em; padding: 15px; width: 220px; text-align: center; border-radius: 16px; border: 4px solid #45475a; background: #11111b; color: white; margin-bottom: 25px; }
        .btn-search { font-size: 2em; padding: 15px 40px; cursor: pointer; background: #89b4fa; color: #11111b; border: none; border-radius: 16px; font-weight: bold; box-shadow: 0 8px 15px rgba(0,0,0,0.3); }
        .btn-search:active { transform: translateY(4px); box-shadow: none; }
        #error-msg { color: #f38ba8; font-weight: bold; font-size: 1.5em; margin-top: 20px; height: 30px; }

        #view-stats { flex: 1; display: none; flex-direction: column; padding: 15px; max-width: 1000px; margin: 0 auto; width: 100%; box-sizing: border-box; overflow-y: auto;}
        
        .stats-header { display: flex; justify-content: space-between; align-items: center; width: 100%; margin-bottom: 15px; }
        .logos-stats { display: flex; gap: 20px; }
        .logos-stats img { height: 60px; border-radius: 8px; }
        
        .btn-retour { align-self: flex-start; font-size: 1.2em; padding: 10px 20px; margin-bottom: 0; cursor: pointer; background: #45475a; color: #cdd6f4; border: none; border-radius: 12px; font-weight: bold; }
        
        .runner-name { font-size: 2.2em; color: #ffffff; font-weight: bold; margin-bottom: 15px; border-bottom: 3px solid #313244; padding-bottom: 10px; text-align: center; display: flex; flex-direction: column; align-items: center; gap: 8px;}
        
        .stats-grid { display: grid; grid-template-columns: repeat(3, 1fr); gap: 10px; margin-bottom: 20px; }
        .stat-card { background: #181825; padding: 15px 5px; border-radius: 12px; border-top: 5px solid #a6e3a1; text-align: center; font-size: 1.1em;}
        .stat-val { font-size: 1.6em; font-weight: bold; color: #f9e2af; margin-top: 8px; }
        
        table { width: 100%; border-collapse: collapse; margin-top: 10px; background: #181825; border-radius: 10px; overflow: hidden; font-size: 1.1em;}
        th, td { padding: 12px 5px; border-bottom: 2px solid #313244; text-align: center;}
        th { background: #313244; color: #ffffff; }
        .status-ok { color: #a6e3a1; font-weight: bold;}
        .status-dnf { color: #f38ba8; font-weight: bold;}

        .best-lap-row { background-color: rgba(166, 227, 161, 0.15) !important; border-left: 6px solid #a6e3a1 !important; }
        .best-lap-row td { color: #a6e3a1 !important; font-weight: bold !important; }

        @media (max-width: 768px) {
            .logos-container { gap: 20px; }
            .logos-container img { height: 80px; } 
            h1 { font-size: 2.2em; margin-bottom: 20px;}
            .stats-header { flex-direction: column; gap: 15px; align-items: flex-start; }
            .logos-stats { width: 100%; justify-content: center; gap: 15px; }
            .logos-stats img { height: 45px; }
            .runner-name { font-size: 1.6em; }
            .stats-grid { grid-template-columns: repeat(2, 1fr); }
            .stat-card { font-size: 0.9em; padding: 10px 2px;}
            .stat-val { font-size: 1.3em; }
            table { font-size: 0.9em; }
            th, td { padding: 8px 2px; }
        }
    </style>
</head>
<body>
    <div id="view-search">
        <div class="logos-container">
            <img src="/static/logo_backyard.png" alt="Logo">
            <img src="/static/logo_tdr.png" alt="Logo TDR">
            <img src="/static/logo_fitness_park.png" alt="Logo Fitness Park">
        </div>
        <h1>Vos Statistiques</h1>
        <input type="number" id="input-dossard" placeholder="Dossard" autocomplete="off">
        <button class="btn-search" onclick="chercher()">Voir mes temps</button>
        <div id="error-msg"></div>
    </div>

    <div id="view-stats">
        <div class="stats-header">
            <button class="btn-retour" onclick="retour()">Retour</button>
            <div class="logos-stats">
                <img src="/static/logo_backyard.png" alt="Logo">
                <img src="/static/logo_tdr.png" alt="Logo TDR">
                <img src="/static/logo_fitness_park.png" alt="Logo Fitness Park">
            </div>
        </div>
        
        <div class="runner-name" id="runner-name"></div>
        <div class="stats-grid">
            <div class="stat-card">Boucles Validées<div class="stat-val" id="val-boucles"></div></div>
            <div class="stat-card">Distance Totale<div class="stat-val" id="val-distance" style="color:#89b4fa;"></div></div>
            <div class="stat-card">Meilleur Tour<div class="stat-val" id="val-fastest"></div></div>
            <div class="stat-card">Temps Moyen<div class="stat-val" id="val-moy"></div></div>
            <div class="stat-card">Allure Moyenne<div class="stat-val" id="val-allure"></div></div>
            <div class="stat-card">Repos Cumulé<div class="stat-val" id="val-repos" style="color:#a6e3a1;"></div></div>
        </div>
        <table>
            <thead><tr><th>N° Boucle</th><th>Temps du tour</th><th>Statut</th></tr></thead>
            <tbody id="table-body"></tbody>
        </table>
    </div>

    <script>
        document.getElementById("input-dossard").addEventListener("keypress", function(event) {
            if (event.key === "Enter") { event.preventDefault(); chercher(); }
        });

        async function chercher() {
            const inputEl = document.getElementById('input-dossard');
            const dossard = inputEl.value.trim();
            const errEl = document.getElementById('error-msg');
            const viewSearch = document.getElementById('view-search');
            const viewStats = document.getElementById('view-stats');
            
            if(!dossard) {
                errEl.innerText = "Veuillez entrer un numéro !";
                return;
            }

            errEl.innerText = "Recherche en cours...";

            const res = await fetch('/api/coureur/' + dossard);
            const data = await res.json();

            if(data.error) {
                errEl.innerText = data.error;
                inputEl.value = ""; 
                inputEl.focus();
            } else {
                errEl.innerText = '';
                viewSearch.style.display = 'none';
                viewStats.style.display = 'flex';
                
                document.getElementById('runner-name').innerHTML = `
                    <div>#${data.dossard} - ${data.nom}</div>
                    <div style="font-size: 0.6em; color: #f9e2af; background: rgba(249, 226, 175, 0.1); padding: 5px 15px; border-radius: 20px; display: inline-block;">${data.rang}e / ${data.total_coureurs}</div>
                `;
                
                document.getElementById('val-boucles').innerText = data.tours_total;
                document.getElementById('val-distance').innerText = data.distance;
                document.getElementById('val-fastest').innerText = data.fastest;
                document.getElementById('val-moy').innerText = data.moyenne;
                document.getElementById('val-allure').innerText = data.allure;
                document.getElementById('val-repos').innerText = data.repos_total;

                document.getElementById('table-body').innerHTML = data.details.map(t => {
                    let isBest = (t.boucle == data.fastest_boucle_num && t.statut.includes('OK'));
                    let rowClass = isBest ? "best-lap-row" : "";
                    let icon = isBest ? "" : "";
                    let statusClass = t.statut.includes('OK') ? 'status-ok' : 'status-dnf';
                    
                    return `<tr class="${rowClass}">
                        <td>B${t.boucle}</td>
                        <td>${t.temps}${icon}</td>
                        <td class="${statusClass}">${t.statut}</td>
                    </tr>`;
                }).reverse().join(''); 
            }
        }

        function retour() {
            document.getElementById('view-stats').style.display = 'none';
            document.getElementById('view-search').style.display = 'flex';
            const inputEl = document.getElementById('input-dossard');
            inputEl.value = "";
            inputEl.focus();
        }
    </script>
</body>
</html>
"""

HTML_PODIUM = """
<!DOCTYPE html>
<html lang="fr">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Podium Backyard Ultra</title>
    <style>
        body { font-family: 'Helvetica', Arial, sans-serif; background-color: #11111b; color: #cdd6f4; text-align: center; margin: 0; padding: 20px 30px; height: 100vh; box-sizing: border-box; display: flex; flex-direction: column; overflow: hidden;}
        
        .header-logos { display: flex; justify-content: center; gap: 50px; margin-bottom: 10px; align-items: center;}
        .header-logos img { height: 130px; border-radius: 12px;}
        
        h1 { font-size: 3.5em; color: #f9e2af; margin-top: 10px; margin-bottom: 20px; text-transform: uppercase; letter-spacing: 3px; text-shadow: 0 4px 10px rgba(249, 226, 175, 0.2); }
        
        .podium-grid { display: grid; grid-template-columns: 1fr 1fr; gap: 20px; max-width: 1300px; margin: 0 auto; flex-grow: 1; width: 100%; min-height: 0;}
        
        .card { background: #181825; border-radius: 16px; padding: 20px; display: flex; flex-direction: column; align-items: center; box-shadow: 0 10px 20px rgba(0,0,0,0.4); overflow: hidden; }
        
        .card-title { font-size: 1.8em; font-weight: bold; padding-bottom: 15px; margin-bottom: 15px; color: #ffffff; text-transform: uppercase; letter-spacing: 2px; width: 100%; border-bottom: 2px solid rgba(255,255,255,0.1); }
        
        .runners-container { display: flex; flex-direction: column; justify-content: center; align-items: center; width: 100%; height: 100%; overflow: hidden;}
        .runner-block { padding: 10px; border-radius: 8px; text-align: center;} 

        .runner-item { font-weight: bold; margin-bottom: 10px; line-height: 1.2; font-size: 2.8em;}
        .runner-detail { color: #a6adc8; font-family: monospace; font-size: 1.5em;}
        
        .card-gold { border-top: 8px solid #f9e2af; }
        .card-gold .runner-item { color: #f9e2af; }
        
        .card-silver { border-top: 8px solid #a6adc8; }
        .card-silver .runner-item { color: #a6adc8; }
        
        .card-bronze { border-top: 8px solid #fab387; }
        .card-bronze .runner-item { color: #fab387; }
        
        .card-blue { border-top: 8px solid #89b4fa; }
        .card-blue .runner-item { color: #89b4fa; }

        @media (max-width: 768px) {
            body { overflow-y: auto; height: auto; }
            .podium-grid { grid-template-columns: 1fr; overflow: visible; }
            .header-logos img { height: 70px; }
            h1 { font-size: 2em; }
            .card-title { font-size: 1.4em; }
            .runner-item { font-size: 2em; }
        }
    </style>
</head>
<body>
    <div class="header-logos">
        <img src="/static/logo_backyard.png" alt="Logo">
        <img src="/static/logo_tdr.png" alt="Logo TDR">
        <img src="/static/logo_fitness_park.png" alt="Logo Fitness Park">
    </div>
    
    <h1>Résultats Officiels</h1>
    
    <div class="podium-grid">
        <div class="card card-gold">
            <div class="card-title">Vainqueur</div>
            <div id="vainqueurs" class="runners-container">Chargement...</div>
        </div>
        
        <div class="card card-silver">
            <div class="card-title">Assist</div>
            <div id="assists" class="runners-container">Chargement...</div>
        </div>
        
        <div class="card card-bronze" style="justify-content: center;">
            <div class="card-title" style="border: none;">Record du Tour</div>
            <div id="record" class="runners-container" style="align-items: center;">Chargement...</div>
        </div>
        
        <div class="card card-blue" style="justify-content: center;">
            <div class="card-title" style="border: none;">Meilleure Moyenne</div>
            <div id="moyenne" class="runners-container" style="align-items: center;">Chargement...</div>
        </div>
    </div>

    <script>
        async function loadPodium() {
            const res = await fetch("/api/podium");
            const data = await res.json();

            if (data.error) {
                document.getElementById('vainqueurs').innerText = data.error;
                return;
            }

            if (data.vainqueurs && data.vainqueurs.length > 0) {
                document.getElementById('vainqueurs').innerHTML = data.vainqueurs.map(v => 
                    `<div class="runner-block"><div class="runner-item">${v.texte}</div><div class="runner-detail">${v.detail}</div></div>`
                ).join('');
            } else {
                document.getElementById('vainqueurs').innerHTML = `<div class="runner-block"><div class="runner-detail" style="font-size: 1.5em;">Aucun</div></div>`;
            }

            if (data.assists && data.assists.length > 0) {
                document.getElementById('assists').innerHTML = data.assists.map(a => 
                    `<div class="runner-block"><div class="runner-item">${a.texte}</div><div class="runner-detail">${a.detail}</div></div>`
                ).join('');
            } else {
                document.getElementById('assists').innerHTML = `<div class="runner-block"><div class="runner-detail" style="font-size: 1.5em;">Aucun</div></div>`;
            }

            if (data.record_tour) {
                document.getElementById('record').innerHTML = `
                    <div class="runner-block">
                        <div class="runner-item">${data.record_tour.runner}</div>
                        <div class="runner-detail">${data.record_tour.detail}</div>
                    </div>`;
            } else {
                document.getElementById('record').innerHTML = `<div class="runner-block"><div class="runner-detail" style="font-size: 1.5em;">Aucun</div></div>`;
            }

            if (data.meilleure_moy) {
                document.getElementById('moyenne').innerHTML = `
                    <div class="runner-block">
                        <div class="runner-item">${data.meilleure_moy.runner}</div>
                        <div class="runner-detail">Moyenne: ${data.meilleure_moy.detail}</div>
                    </div>`;
            } else {
                document.getElementById('moyenne').innerHTML = `<div class="runner-block"><div class="runner-detail" style="font-size: 1.5em;">Aucun</div></div>`;
            }
        }
        
        setInterval(loadPodium, 10000);
        loadPodium();
    </script>
</body>
</html>
"""

HTML_CLASSEMENT = """
<!DOCTYPE html>
<html lang="fr">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Classement Général - Backyard Ultra</title>
    <style>
        body { font-family: 'Helvetica', Arial, sans-serif; background-color: #1e1e2e; color: #cdd6f4; text-align: center; margin: 0; padding: 30px; }
        .header-logos { display: flex; justify-content: center; gap: 50px; margin-bottom: 20px; align-items: center;}
        .header-logos img { height: 100px; border-radius: 12px;}
        h1 { font-size: 3em; color: #f9e2af; margin-top: 0; margin-bottom: 10px; text-transform: uppercase; letter-spacing: 2px; }
        .btn-export { font-size: 1.2em; padding: 12px 25px; margin-bottom: 30px; cursor: pointer; background: #a6e3a1; color: #11111b; border: none; border-radius: 12px; font-weight: bold; text-decoration: none; display: inline-block; box-shadow: 0 4px 10px rgba(0,0,0,0.3);}
        .btn-export:hover { filter: brightness(1.1); }
        .table-wrapper { width: 100%; max-width: 1200px; margin: 0 auto; background: #181825; border-radius: 12px; overflow-x: auto; box-shadow: 0 10px 20px rgba(0,0,0,0.3); }
        table { width: 100%; border-collapse: collapse; font-size: 1.2em;}
        th, td { padding: 15px; border-bottom: 1px solid #313244; text-align: center; white-space: nowrap;}
        th { background: #313244; color: #ffffff; position: sticky; top: 0; z-index: 10; text-transform: uppercase; letter-spacing: 1px;}
        tr:hover { background-color: rgba(255,255,255,0.05); }
        .rang-1 { color: #f9e2af; font-weight: bold; font-size: 1.3em;}
        .rang-2 { color: #a6adc8; font-weight: bold; font-size: 1.2em;}
        .rang-3 { color: #fab387; font-weight: bold; font-size: 1.2em;}
        .col-tours { color: #a6e3a1; font-weight: bold; font-size: 1.2em;}
        .col-distance { color: #89b4fa; }

        @media (max-width: 768px) {
            body { padding: 10px; }
            .header-logos img { height: 60px; }
            h1 { font-size: 2em; }
            table { font-size: 0.9em; }
            th, td { padding: 10px; }
        }
    </style>
</head>
<body>
    <div class="header-logos">
        <img src="/static/logo_backyard.png" alt="Logo">
        <img src="/static/logo_tdr.png" alt="Logo TDR">
        <img src="/static/logo_fitness_park.png" alt="Logo Fitness Park">
    </div>
    
    <h1>Classement Général</h1>
    
    <a href="/api/export_excel" class="btn-export">Exporter tous les résultats (Excel)</a>
    <a href="/api/download_statuts" class="btn-export" style="background: #89b4fa; margin-left: 15px;">Télécharger le statut des éliminations (Excel)</a>
    
    <div class="table-wrapper">
        <table>
            <thead>
                <tr>
                    <th>Rang</th>
                    <th>Dossard</th>
                    <th style="text-align: left;">Nom</th>
                    <th>Boucles</th>
                    <th>Distance</th>
                    <th>Meilleur Tour</th>
                    <th>Moyenne</th>
                </tr>
            </thead>
            <tbody id="table-body">
                <tr><td colspan="7">Chargement en cours...</td></tr>
            </tbody>
        </table>
    </div>

    <script>
        async function loadClassement() {
            const res = await fetch("/api/classement");
            const data = await res.json();
            
            document.getElementById('table-body').innerHTML = data.classement.map(c => {
                let rankClass = '';
                let rankTrophy = '';
                
                if (c.rang === 1) { rankClass = 'rang-1'; rankTrophy = ''; }
                else if (c.rang === 2) { rankClass = 'rang-2'; rankTrophy = ''; }
                else if (c.rang === 3) { rankClass = 'rang-3'; rankTrophy = ''; }

                return `<tr>
                    <td class="${rankClass}">${rankTrophy}${c.rang}</td>
                    <td style="font-weight: bold;">#${c.dossard}</td>
                    <td style="text-align: left; font-weight: bold; color: #ffffff;">${c.nom}</td>
                    <td class="col-tours">${c.tours}</td>
                    <td class="col-distance">${c.distance}</td>
                    <td style="font-family: monospace;">${c.fastest}</td>
                    <td style="font-family: monospace;">${c.moyenne}</td>
                </tr>`;
            }).join('');
        }
        
        setInterval(loadClassement, 15000);
        loadClassement();
    </script>
</body>
</html>
"""

HTML_SPEAKER = """
<!DOCTYPE html>
<html lang="fr">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Écran Speaker - Dashboard Interactif</title>
    <style>
        body { font-family: 'Helvetica', Arial, sans-serif; background-color: #11111b; color: #ffffff; text-align: center; margin: 0; padding: 15px; min-height: 100vh; height: auto; box-sizing: border-box; display: flex; flex-direction: column; overflow-y: auto; }
        .header { display: flex; justify-content: space-between; align-items: center; border-bottom: 4px solid #313244; padding-bottom: 10px; margin-bottom: 15px;}
        .titre-boucle { font-size: 3em; font-weight: bold; color: #f9e2af; margin: 0;}
        .clock { font-size: 4em; font-weight: bold; color: #a6e3a1; margin: 0; line-height: 1;}
        .danger { color: #f38ba8 !important; animation: blinker 1s linear infinite; }
        @keyframes blinker { 50% { opacity: 0.5; } }
        .zones-container { display: flex; flex-direction: column; flex-grow: 1; gap: 20px; height: auto; overflow: visible; }
        .zone { display: flex; flex-direction: column; background: #181825; border-radius: 12px; padding: 15px; }
        .zone-title { font-size: 1.5em; font-weight: bold; text-align: left; margin-bottom: 15px; border-bottom: 2px solid; padding-bottom: 5px;}
        .grid-cards { display: grid; grid-template-columns: repeat(auto-fill, minmax(200px, 1fr)); gap: 10px; }
        .runner-card { border-radius: 8px; padding: 12px; font-size: 1.2em; font-weight: bold; cursor: pointer; transition: transform 0.1s; display: flex; flex-direction: column; justify-content: center; align-items: center;}
        .runner-card:hover { transform: scale(1.02); filter: brightness(1.2);}
        .zone-encours { border: 2px solid #f38ba8;} 
        .zone-encours .zone-title { color: #f38ba8; border-color: #f38ba8;}
        .card-encours { background: rgba(243, 139, 168, 0.15); border: 1px solid #f38ba8; color: #f38ba8;}
        .zone-arrive { border: 2px solid #a6e3a1;}
        .zone-arrive .zone-title { color: #a6e3a1; border-color: #a6e3a1;}
        .card-arrive { background: rgba(166, 227, 161, 0.15); border: 1px solid #a6e3a1; color: #a6e3a1; font-size: 1.1em; padding: 10px;}
        .badge-record { color: #11111b; background-color: #f9e2af; font-size: 0.75em; padding: 3px 8px; border-radius: 4px; margin-top: 6px; animation: blinker 1s linear infinite; width: 90%;}
        .modal { display: none; position: fixed; z-index: 1000; left: 0; top: 0; width: 100%; height: 100%; background-color: rgba(0,0,0,0.8); justify-content: center; align-items: center; padding: 10px; box-sizing: border-box;}
        .modal-content { background-color: #1e1e2e; padding: 20px; border-radius: 16px; width: 100%; max-width: 800px; border: 4px solid #89b4fa; position: relative;}
        .close-btn { color: #f38ba8; float: right; font-size: 2em; font-weight: bold; cursor: pointer; line-height: 0.5;}
        .modal-name { font-size: 2em; font-weight: bold; margin-bottom: 20px; color: #ffffff;}
        .stats-row { display: flex; justify-content: space-between; background: #181825; padding: 15px; border-radius: 12px; margin-bottom: 20px;}
        .stat-item { text-align: center; flex: 1;}
        .stat-label { color: #a6adc8; font-size: 0.9em; }
        .stat-value { color: #89b4fa; font-size: 1.5em; font-weight: bold; margin-top: 5px;}
        .modal table { width: 100%; border-collapse: collapse; font-size: 1.1em;}
        .modal th, .modal td { padding: 8px; border-bottom: 1px solid #313244;}
        .modal th { color: #a6adc8; }
        .best-lap-row { background-color: rgba(166, 227, 161, 0.15) !important; border-left: 6px solid #a6e3a1 !important; color: #a6e3a1 !important; font-weight: bold;}

        @media (max-width: 768px) {
            .titre-boucle { font-size: 2em; }
            .clock { font-size: 3em; }
            .stats-row { flex-wrap: wrap; gap: 10px;}
            .stat-item { min-width: 45%; }
            .grid-cards { grid-template-columns: repeat(2, 1fr); gap: 8px; }
            .runner-card { font-size: 1em; padding: 8px; }
        }
    </style>
</head>
<body>
    <div class="header">
        <h1 class="titre-boucle" id="loop-info">EN ATTENTE</h1>
        <div class="clock" id="countdown">--:--</div>
    </div>
    
    <div class="zones-container">
        <div class="zone zone-encours">
            <div class="zone-title" id="title-encours">Sur le parcours (--)</div>
            <div class="grid-cards" id="list-encours"></div>
        </div>
        
        <div class="zone zone-arrive">
            <div class="zone-title" id="title-arrive">Viennent d'arriver (--)</div>
            <div class="grid-cards" id="list-arrive"></div>
        </div>
    </div>

    <div id="runnerModal" class="modal">
        <div class="modal-content">
            <span class="close-btn" onclick="closeModal()">&times;</span>
            <div class="modal-name" id="modal-name">#-- Nom du Coureur</div>
            
            <div class="stats-row">
                <div class="stat-item"><div class="stat-label">Boucles</div><div class="stat-value" id="modal-boucles">--</div></div>
                <div class="stat-item"><div class="stat-label">Meilleur</div><div class="stat-value" style="color:#a6e3a1;" id="modal-fastest">--:--</div></div>
                <div class="stat-item"><div class="stat-label">Moyenne</div><div class="stat-value" id="modal-moy">--:--</div></div>
                <div class="stat-item"><div class="stat-label">Rang</div><div class="stat-value" style="color:#f9e2af;" id="modal-rang">--</div></div>
            </div>
            
            <div style="max-height: 250px; overflow-y: auto;">
                <table>
                    <thead><tr><th>N° Boucle</th><th>Temps</th><th>Statut</th></tr></thead>
                    <tbody id="modal-table"></tbody>
                </table>
            </div>
        </div>
    </div>

    <script>
        let currentLoop = 1;

        async function fetchState() {
            const res = await fetch("/api/state");
            const data = await res.json();
            
            if (data.started) {
                currentLoop = data.boucle;
                document.getElementById("loop-info").innerText = `BOUCLE ${data.boucle}`;
                
                let sec = data.temps_restant;
                let clockEl = document.getElementById("countdown");
                if (sec > 0) {
                    let m = Math.floor(sec / 60).toString().padStart(2, '0');
                    let s = (sec % 60).toString().padStart(2, '0');
                    clockEl.innerText = `${m}:${s}`;
                    clockEl.className = sec < 180 ? "clock danger" : "clock";
                } else {
                    clockEl.innerText = "00:00";
                    clockEl.className = "clock danger";
                }

                let enCourse = data.runners_all.filter(r => r.status === 'encourse');
                let arrives = data.runners_all.filter(r => r.status === 'arrive');
                
                arrives.sort((a, b) => a.last_loop_sec - b.last_loop_sec);
                
                document.getElementById("title-encours").innerText = `Sur le parcours (${enCourse.length})`;
                document.getElementById("title-arrive").innerText = `Arrivés (Boucle ${currentLoop}) - (${arrives.length})`;
                
                document.getElementById("list-encours").innerHTML = enCourse.map(r => 
                    `<div class="runner-card card-encours" onclick="openModal('${r.dossard}')">#${r.dossard} - ${r.nom}</div>`
                ).join('');
                
                document.getElementById("list-arrive").innerHTML = arrives.map(r => {
                    let badge = r.just_beat_record ? `<div class="badge-record">RECORD BATTU !</div>` : ``;
                    return `<div class="runner-card card-arrive" onclick="openModal('${r.dossard}')">
                        <span>#${r.dossard} - ${r.nom}</span>
                        <span style="font-size:0.8em; font-weight:normal; margin-top:3px;">${r.last_loop}</span>
                        ${badge}
                    </div>`;
                }).join('');
            }
        }

        async function openModal(dossard) {
            const res = await fetch('/api/coureur/' + dossard);
            const data = await res.json();
            
            if(!data.error) {
                document.getElementById('modal-name').innerText = `#${data.dossard} - ${data.nom}`;
                document.getElementById('modal-boucles').innerText = data.tours_total;
                document.getElementById('modal-fastest').innerText = data.fastest;
                document.getElementById('modal-moy').innerText = data.moyenne;
                document.getElementById('modal-rang').innerText = `${data.rang}e`;
                
                document.getElementById('modal-table').innerHTML = data.details.map((t, index) => {
                    let isBest = (t.boucle == data.fastest_boucle_num && t.statut.includes('OK'));
                    let rowClass = isBest ? "best-lap-row" : "";
                    let icon = isBest ? "" : "";
                    let statusColor = t.statut.includes('OK') ? '#a6e3a1' : '#f38ba8';
                    
                    return `<tr class="${rowClass}">
                        <td>B${t.boucle}</td>
                        <td style="color:${isBest ? '#a6e3a1' : 'inherit'}; font-weight:${isBest ? 'bold' : 'normal'};">${t.temps}${icon}</td>
                        <td style="color:${statusColor}; font-weight:bold;">${t.statut}</td>
                    </tr>`;
                }).reverse().join(''); 

                document.getElementById('runnerModal').style.display = "flex";
            }
        }

        function closeModal() {
            document.getElementById('runnerModal').style.display = "none";
        }
        
        window.onclick = function(event) {
            let modal = document.getElementById('runnerModal');
            if (event.target == modal) { closeModal(); }
        }

        setInterval(fetchState, 1000);
        fetchState();
    </script>
</body>
</html>
"""

@app.route("/")
def index():
    return render_template_string(HTML_MAIN)

@app.route("/tv")
def tv_index():
    return render_template_string(HTML_TV)

@app.route("/coureur")
def coureur_index():
    return render_template_string(HTML_COUREUR)

@app.route("/podium")
def podium_index():
    return render_template_string(HTML_PODIUM)

@app.route("/classement")
def classement_index():
    return render_template_string(HTML_CLASSEMENT)

@app.route("/speaker")
def speaker_index():
    return render_template_string(HTML_SPEAKER)

@app.route("/roulette")
def roulette_index():
    return render_template_string(HTML_ROULETTE)

if __name__ == "__main__":
    app.run(host="0.0.0.0", port=8000, debug=True)